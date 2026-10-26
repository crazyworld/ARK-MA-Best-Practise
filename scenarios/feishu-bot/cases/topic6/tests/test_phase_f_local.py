#!/usr/bin/env python3
"""Phase F + HC3 本地端到端验证。

在 212 机器执行，模拟沙箱中飞书 API 的完整调用链路。
验证：凭证→创建文件夹→创建在线表格→写入数据→清理。

用法：
  source ~/.arkagent/cases/topic6/config.env
  cd /home/bmc/work/ARK-MA-Best-Practise/scenarios/feishu-bot
  python cases/topic6/tests/test_phase_f_local.py
"""
from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

_CASE_DIR = Path(__file__).resolve().parents[1]
_FEISHU_BOT_DIR = _CASE_DIR.parents[1]
for p in (str(_FEISHU_BOT_DIR), str(_CASE_DIR)):
    if p not in sys.path:
        sys.path.insert(0, p)

import requests

APP_ID = os.environ.get("LARK_APP_ID") or os.environ.get("FEISHU_APP_ID") or ""
APP_SECRET = os.environ.get("LARK_APP_SECRET") or os.environ.get("FEISHU_APP_SECRET") or ""
FOLDER_TOKEN = os.environ.get("TOPIC6_DOCS_FOLDER_TOKEN") or "CHz0fnO5hlJAVddjiWBchhvlnSg"

G = "\033[92m"
R = "\033[91m"
NC = "\033[0m"
errors: list[str] = []


def ok(msg): print(f"{G}✅ {msg}{NC}")
def fail(msg): print(f"{R}❌ {msg}{NC}"); errors.append(msg)


def cleanup(headers, folder="", sheet_token=""):
    try:
        if sheet_token:
            requests.delete(f"https://open.feishu.cn/open-apis/drive/v1/files/{sheet_token}?type=sheet", headers=headers, timeout=10)
        if folder:
            requests.delete(f"https://open.feishu.cn/open-apis/drive/v1/files/{folder}?type=folder", headers=headers, timeout=10)
    except Exception:
        pass


def main() -> int:
    print("===== Phase F + HC3 本地验证 =====")
    print(f"APP_ID: {APP_ID[:8]}..." if APP_ID else "APP_ID: EMPTY")
    print(f"FOLDER_TOKEN: {FOLDER_TOKEN}\n")

    if not APP_ID or not APP_SECRET:
        fail("飞书凭证未配置"); return 1

    # 1. token
    print("--- 1/5 获取 tenant_access_token ---")
    resp = requests.post("https://open.feishu.cn/open-apis/auth/v3/tenant_access_token/internal",
        json={"app_id": APP_ID, "app_secret": APP_SECRET}, timeout=10)
    token = resp.json().get("tenant_access_token", "")
    if not token:
        fail(f"token 失败: {resp.json()}"); return 1
    ok(f"token OK (len={len(token)})")
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json; charset=utf-8"}

    # 2. 创建文件夹
    print("--- 2/5 创建测试文件夹 ---")
    resp = requests.post("https://open.feishu.cn/open-apis/drive/v1/files/create_folder",
        headers=headers, json={"name": "_phase_f_e2e_test", "folder_token": FOLDER_TOKEN}, timeout=30)
    data = resp.json()
    folder = (data.get("data") or {}).get("token") or (data.get("data") or {}).get("folder_token", "")
    if data.get("code") == 0 and folder:
        ok(f"create_folder OK token={folder}")
    else:
        fail(f"create_folder 失败: {data}"); return 1

    # 3. 创建在线表格 + 写入数据
    print("--- 3/5 创建在线表格并写入数据 ---")
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["row_id", "platform", "title", "score"])
    ws.append([1, "微博", "测试热搜", 95])
    ws.append([2, "抖音", "测试话题", 88])
    tf = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(tf.name)

    resp = requests.post("https://open.feishu.cn/open-apis/sheets/v3/spreadsheets",
        headers=headers, json={"title": "phase_f_test", "folder_token": folder}, timeout=30)
    data = resp.json()
    sheet_token = (data.get("data") or {}).get("spreadsheet", {}).get("spreadsheet_token", "")
    if not (data.get("code") == 0 and sheet_token):
        fail(f"创建在线表格失败: {data}")
        os.unlink(tf.name); cleanup(headers, folder=folder); return 1

    # 获取 sheet_id
    resp = requests.get(f"https://open.feishu.cn/open-apis/sheets/v3/spreadsheets/{sheet_token}/sheets/query",
        headers=headers, timeout=10)
    sheet_id = (resp.json().get("data") or {}).get("sheets", [{}])[0].get("sheet_id", "Sheet1")

    # 读 xlsx 写入
    wb2 = openpyxl.load_workbook(tf.name, read_only=True, data_only=True)
    rows = [[v if v is not None else "" for v in row] for row in wb2.active.iter_rows(values_only=True)]
    wb2.close()
    os.unlink(tf.name)

    num_cols = max(len(r) for r in rows)
    end_col, n = "", num_cols
    while n > 0:
        n -= 1; end_col = chr(65 + n % 26) + end_col; n //= 26

    resp = requests.put(f"https://open.feishu.cn/open-apis/sheets/v2/spreadsheets/{sheet_token}/values",
        headers=headers,
        json={"valueRange": {"range": f"{sheet_id}!A1:{end_col}{len(rows)}", "values": rows}},
        timeout=30)
    write_ok = resp.json().get("code") == 0
    url = f"https://bluefocus.feishu.cn/sheets/{sheet_token}"
    if write_ok:
        ok(f"在线表格 OK → {url}")
    else:
        fail(f"写入数据失败: {resp.json()}")

    # 4. environment.json 检查
    print("--- 4/5 environment.json 完整性 ---")
    env_data = json.loads((_CASE_DIR / "ma-resources" / "environment.json").read_text())
    env = env_data["config"]["env"]
    required = ["LARK_APP_ID", "LARK_APP_SECRET", "LARKSUITE_CLI_APP_ID"]
    forbidden = ["LARKSUITE_CLI_APP_SECRET", "LARKSUITE_CLI_STRICT_MODE"]  # STRICT_MODE=bot 会拦 Phase H --as user
    missing = [k for k in required if k not in env] + [f'禁止:{k}' for k in forbidden if k in env]
    if missing:
        fail(f"缺少: {missing}")
    else:
        ok("environment.json 完整")

    # 5. 清理
    print("--- 5/5 清理 ---")
    cleanup(headers, folder=folder, sheet_token=sheet_token)
    ok("清理完成")

    print()
    if errors:
        print(f"{R}===== {len(errors)} 项失败 ====={NC}")
        for e in errors: print(f"  • {e}")
        return 1
    print(f"{G}===== 全部通过 ====={NC}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
